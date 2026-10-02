"""기존 '전장 업무전달 리스트' 레이아웃으로 엑셀 내보내기."""

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from app.rules import OVERDUE, DONE, status_of

HEADER_FILL = PatternFill("solid", fgColor="FFE699")  # 기존 노란 머리글
OVERDUE_FONT = Font(color="C00000")
DONE_FONT = Font(color="808080")
THIN = Side(style="thin", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

# (머리글, 열 너비) — 열 순서는 지시서 5장-2 기준
COLUMNS = [
    ("프로젝트 차종", 14),
    ("작성날짜", 11),
    ("목표", 11),
    ("완료", 11),
    ("업무내용", 60),
    ("담당", 10),
    ("진행내용", 40),
    ("파일경로", 25),
    ("비고", 15),
]


def target_cell(task: dict) -> str:
    """목표 변경 이력이 있으면 '최초목표↵현재목표'."""
    history = [h["date"] for h in task["target_history"]]
    if len(history) > 1 and history[0] != history[-1]:
        return f"{history[0]}\n{history[-1]}"
    return task["target"] or ""


def owner_cell(task: dict) -> str:
    """'주담당↵협업자' (협업자가 여러 명이면 한 줄에 하나씩)."""
    return "\n".join([task["owner"] or "", *task["collaborators"]]).strip()


def progress_cell(task: dict) -> str:
    return "\n".join(f"{c['ts'][5:10]} {c['by']}: {c['text']}" for c in task["comments"])


def build_workbook(tasks: list[dict], today: date) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "업무전달리스트"

    # 1행: 제목 + 날짜
    ws["A1"] = "■ 전장 업무전달 리스트"
    ws["A1"].font = Font(bold=True, size=14)
    ws["I1"] = f"{today.isoformat()} 기준"
    ws["I1"].alignment = Alignment(horizontal="right")

    # 2~3행: '일정'이 작성날짜·목표·완료 위에 병합, 나머지 머리글은 세로 병합
    for col, (title, width) in enumerate(COLUMNS, start=1):
        letter = ws.cell(row=3, column=col).column_letter
        ws.column_dimensions[letter].width = width
        if 2 <= col <= 4:
            ws.cell(row=3, column=col, value=title)
        else:
            ws.cell(row=2, column=col, value=title)
            ws.merge_cells(start_row=2, start_column=col, end_row=3, end_column=col)
    ws.cell(row=2, column=2, value="일정")
    ws.merge_cells("B2:D2")
    for row in ws.iter_rows(min_row=2, max_row=3, max_col=len(COLUMNS)):
        for cell in row:
            cell.fill, cell.border, cell.alignment = HEADER_FILL, BORDER, CENTER
            cell.font = Font(bold=True)

    # 4행부터 데이터
    for r, task in enumerate(tasks, start=4):
        values = [
            task["project"] or "",
            task["created_date"] or "",
            target_cell(task),
            task["done_date"] or "",
            task["content"] or "",
            owner_cell(task),
            progress_cell(task),
            task["file_path"] or "",
            task["note"] or "",
        ]
        status = status_of(task, today)
        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.border, cell.alignment = BORDER, WRAP_TOP
            if status == OVERDUE:
                cell.font = OVERDUE_FONT
            elif status == DONE:
                cell.font = DONE_FONT

    ws.freeze_panes = "A4"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
