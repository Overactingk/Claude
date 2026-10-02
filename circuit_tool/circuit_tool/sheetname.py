"""시트 파일명 해석 (요구사항 4.3).

69001-KV405_001004_(CE_BOX)_R0.dwg  → 도면 69001, 차종 KV405, 1/4장, CE_BOX, R0
69002-KV405_(AUX_FUSE_BOX)_R0.dwg   → 장번호 없음
69032-KV405_001002(NACEKO_SW)_R0    → '_' 누락 형식도 허용
"""
from __future__ import annotations

import re
from pathlib import Path

from .model import Issue, Sheet

_PAT = re.compile(
    r"^(?P<dno>\d+)-(?P<veh>[A-Za-z0-9]+)_"
    r"(?:(?P<sno>\d{3})(?P<stot>\d{3})_?)?"
    r"\((?P<name>[^)]*)\)_"
    r"(?P<rev>R\d+)$",
    re.IGNORECASE,
)


def parse_sheet_name(path: str | Path) -> Sheet | None:
    stem = Path(path).stem
    m = _PAT.match(stem)
    if not m:
        return None
    return Sheet(
        id=stem,
        drawing_no=m["dno"],
        vehicle=m["veh"],
        sheet_no=m["sno"],
        sheet_total=m["stot"],
        name=m["name"],
        rev=m["rev"].upper(),
        source=str(path),
    )


def rev_num(rev: str) -> int:
    return int(rev[1:])


def select_latest(paths: list[Path]) -> tuple[list[tuple[Path, Sheet]], list[Issue]]:
    """같은 도면번호·장번호는 최신 리비전만 남긴다. 해석 불가 파일명은 오류."""
    issues: list[Issue] = []
    best: dict[tuple[str, str | None], tuple[Path, Sheet]] = {}
    for p in paths:
        s = parse_sheet_name(p)
        if s is None:
            issues.append(Issue("파일명 해석 불가", p.stem, p.name, "파일명 규칙과 다름 → 가져오지 않음"))
            continue
        key = (s.drawing_no, s.sheet_no)
        old = best.get(key)
        if old is None or rev_num(s.rev) > rev_num(old[1].rev):
            if old is not None:
                issues.append(Issue("구 리비전 제외", old[1].id, old[0].name, f"{s.rev} 사용"))
            best[key] = (p, s)
        else:
            issues.append(Issue("구 리비전 제외", s.id, p.name, f"{old[1].rev} 사용"))
    return sorted(best.values(), key=lambda t: t[1].id), issues


def check_sheet_list(xlsx: str | Path, sheets: list[Sheet],
                     pattern: str = r"(?<!\d)69\d{3}(?!\d)") -> list[Issue]:
    """회로도 리스트 엑셀의 도면번호와 대조해 누락 시트를 보고한다.

    엑셀 양식을 알 수 없으므로 모든 셀에서 도면번호 패턴을 찾는다 (패턴은 변경 가능).
    """
    from openpyxl import load_workbook

    wb = load_workbook(xlsx, read_only=True, data_only=True)
    listed: set[str] = set()
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if v is not None:
                    listed.update(re.findall(pattern, str(v)))
    have = {s.drawing_no for s in sheets}
    issues = [Issue("누락 시트", dno, Path(xlsx).name, "회로도 리스트에 있으나 가져온 도면에 없음")
              for dno in sorted(listed - have)]
    issues += [Issue("리스트 외 시트", dno, Path(xlsx).name, "가져온 도면이 회로도 리스트에 없음")
               for dno in sorted(have - listed)]
    return issues
