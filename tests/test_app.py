"""핵심 동작 검증: 지연 판정, 동시 사용(공유폴더), 이관, 엑셀, 화면 흐름."""

import io
import json
from html import unescape
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app import rules
from app.export import build_workbook
from app.main import create_app, merge_changes
from app.store import Store
from scripts.import_seed import import_seed

SEED = Path(__file__).resolve().parent.parent / "seed_tasks.json"
TODAY = date(2026, 10, 2)


# ---------- 상태 판정 ----------

@pytest.mark.parametrize(
    ("target", "done", "expected"),
    [
        (date(2026, 10, 1), None, rules.OVERDUE),  # 하루 지남
        (date(2026, 10, 2), None, rules.DUE_SOON),  # D-day
        (date(2026, 10, 5), None, rules.DUE_SOON),  # D-3 (경계)
        (date(2026, 10, 6), None, rules.IN_PROGRESS),  # D-4
        (date(2026, 9, 1), date(2026, 9, 5), rules.DONE),  # 늦게라도 완료면 완료
        (None, None, rules.IN_PROGRESS),  # 목표 없음
    ],
)
def test_task_status(target, done, expected):
    assert rules.task_status(target, done, TODAY) == expected


# ---------- 공유폴더 동시 사용 ----------

def test_two_writers_see_each_other(tmp_path):
    a = Store(tmp_path, writer_name="이가람C_pc1")
    b = Store(tmp_path, writer_name="이재욱M_pc2")
    tid = a.create_task("이가람C", {"content": "회로도", "owner": "이가람C"})
    b.refresh()
    b.update_task("이재욱M", tid, {"note": "B가 수정"})  # B는 자기 파일에만 씀
    a.refresh()
    assert a.tasks[tid]["note"] == "B가 수정"
    assert sorted(p.name for p in (tmp_path / "events").iterdir()) == [
        "이가람C_pc1.jsonl",
        "이재욱M_pc2.jsonl",
    ]


def test_partial_line_is_read_later(tmp_path):
    """다른 PC가 아직 쓰는 중인 줄(줄바꿈 없음)은 건너뛰고 다음에 읽는다."""
    store = Store(tmp_path, writer_name="me_1")
    tid = store.create_task("me", {"content": "x", "owner": "me"})
    other = tmp_path / "events" / "other_2.jsonl"
    line = json.dumps({"id": "z", "ts": "2999-01-01T00:00:00.000000", "by": "other",
                       "type": "task_set", "task": tid, "fields": {"note": "늦게"}}, ensure_ascii=False)
    other.write_text(line, encoding="utf-8")  # 줄바꿈 없음
    store.refresh()
    assert store.tasks[tid]["note"] is None
    other.write_text(line + "\n", encoding="utf-8")
    store.refresh()
    assert store.tasks[tid]["note"] == "늦게"


def test_same_timestamp_keeps_file_order(tmp_path, monkeypatch):
    """Windows 시계 해상도로 같은 시각이 찍혀도 생성 → 수정 순서가 유지돼야 한다."""
    import app.store as store_module

    fixed = store_module.datetime(2026, 10, 2, 9, 0, 0)

    class FrozenDatetime(store_module.datetime):
        @classmethod
        def now(cls, tz=None):
            return fixed

    monkeypatch.setattr(store_module, "datetime", FrozenDatetime)
    store = Store(tmp_path, writer_name="me_1")
    tid = store.create_task("me", {"content": "x", "owner": "me", "target": "2026-10-10"})
    store.update_task("me", tid, {"target": "2026-10-20"})
    store.update_task("me", tid, {"target": "2026-10-30"})
    fresh = Store(tmp_path)
    fresh.refresh()
    assert [h["date"] for h in fresh.tasks[tid]["target_history"]] == [
        "2026-10-10", "2026-10-20", "2026-10-30",
    ]


def test_merge_changes_field_level():
    """다른 필드 동시 수정은 둘 다 반영, 같은 필드는 충돌로 저장 안 함."""
    orig = {"note": "a", "target": "2026-10-10"}
    current = {"note": "남이 바꿈", "target": "2026-10-10"}
    mine = {"note": "내가 바꿈", "target": "2026-10-20"}
    changes, conflicts = merge_changes(current, orig, mine)
    assert changes == {"target": "2026-10-20"}
    assert conflicts == ["note"]


# ---------- 이관 ----------

@pytest.fixture
def seeded(tmp_path) -> Store:
    import_seed(SEED, tmp_path)
    store = Store(tmp_path)
    store.refresh()
    return store


def test_seed_import_keeps_known_issues(seeded):
    tasks = list(seeded.tasks.values())
    assert len(tasks) == 47
    issues = [i for t in tasks for i in rules.input_issues(t, TODAY)]
    # 지시서 9장의 알려진 이슈 건수 그대로 (임의 수정 없음)
    assert issues.count("작성날짜 없음") == 19
    assert issues.count("목표 없음") == 1
    assert issues.count("목표가 작성일보다 앞") == 1
    assert issues.count("프로젝트 없음") == 1
    # 미완료 46건 중 목표 없는 No.28은 판정 불가 → 45건
    assert issues.count("지연") == 45


def test_seed_target_history_and_people(seeded):
    first = next(t for t in seeded.tasks.values() if t["no"] == 1)
    assert [h["date"] for h in first["target_history"]] == ["2026-03-18", "2026-06-19"]
    assert first["target"] == "2026-06-19"
    six = next(t for t in seeded.tasks.values() if t["no"] == 6)
    assert six["owner"] == "이가람C" and six["collaborators"] == ["나대경M"]
    assert seeded.users["나대경M"]["is_admin"] is True
    assert seeded.users["전장"]["can_login"] is False


def test_seed_refuses_double_import(tmp_path):
    import_seed(SEED, tmp_path)
    with pytest.raises(SystemExit):
        import_seed(SEED, tmp_path)


# ---------- 엑셀 ----------

def test_excel_layout(seeded):
    tasks = sorted(seeded.tasks.values(), key=lambda t: t["no"])
    ws = load_workbook(io.BytesIO(build_workbook(tasks, TODAY))).active
    assert ws["A1"].value == "■ 전장 업무전달 리스트"
    assert ws["B2"].value == "일정" and "B2:D2" in [str(r) for r in ws.merged_cells.ranges]
    assert [ws.cell(row=3, column=c).value for c in (2, 3, 4)] == ["작성날짜", "목표", "완료"]
    assert ws["A2"].fill.fgColor.rgb.endswith("FFE699")
    assert ws["C4"].value == "2026-03-18\n2026-06-19"  # No.1: 최초목표↵현재목표
    assert ws["F9"].value == "이가람C\n나대경M"  # No.6: 주담당↵협업자


# ---------- 화면 흐름 ----------

@pytest.fixture
def client_for(tmp_path):
    import_seed(SEED, tmp_path)

    def make(name: str, pc: str):
        local = tmp_path / f"local_{pc}.json"
        local.write_text(json.dumps({"me": name, "install_id": pc}), encoding="utf-8")
        app = create_app(tmp_path, local, today=lambda: TODAY)
        app.testing = True
        return app.test_client()

    return make


def _task_id(client, no: int) -> str:
    html = client.get("/").get_data(as_text=True)
    marker = f'">{no}</a>'
    return html[: html.index(marker)].rsplit("/task/", 1)[1]


def test_pages_render(client_for):
    c = client_for("나대경M", "pc1")
    for url in ["/", "/dashboard", "/checks", "/admin", "/task/new", "/export.xlsx"]:
        assert c.get(url).status_code == 200, url
    assert "⚠ 지연" in c.get("/").get_data(as_text=True)


def test_first_run_asks_for_name(tmp_path):
    import_seed(SEED, tmp_path)
    app = create_app(tmp_path, tmp_path / "local.json", today=lambda: TODAY)
    c = app.test_client()
    assert c.get("/").status_code == 302
    assert c.post("/setup", data={"name": "이가람C"}).status_code == 302
    assert c.get("/").status_code == 200
    assert json.loads((tmp_path / "local.json").read_text(encoding="utf-8"))["me"] == "이가람C"


def test_edit_permission(client_for):
    owner = client_for("이가람C", "pc1")
    other = client_for("박삼용C", "pc2")
    tid = _task_id(owner, 1)  # 이가람C 담당
    assert owner.get(f"/task/{tid}/edit").status_code == 200
    assert other.get(f"/task/{tid}/edit").status_code == 403
    assert other.post(f"/task/{tid}/delete").status_code == 403  # 삭제는 관리자만


def test_concurrent_edit_through_ui(client_for):
    """두 PC가 같은 업무를 동시에 열고 저장: 다른 필드는 둘 다 반영, 같은 필드는 경고."""
    a = client_for("이가람C", "pc1")
    b = client_for("나대경M", "pc2")
    tid = _task_id(a, 1)

    def form_of(client):
        html = client.get(f"/task/{tid}/edit").get_data(as_text=True)
        orig_raw = html.split('name="orig" value="', 1)[1].split('"', 1)[0]
        orig = json.loads(unescape(orig_raw))
        data = {k: (v or "") for k, v in orig.items() if k != "collaborators"}
        return orig, data

    orig_a, data_a = form_of(a)
    orig_b, data_b = form_of(b)
    data_a.update(note="A 비고", orig=json.dumps(orig_a))
    data_b.update(note="B 비고", done_date="2026-10-01", orig=json.dumps(orig_b))
    a.post(f"/task/{tid}/edit", data=data_a)
    resp = b.post(f"/task/{tid}/edit", data=data_b, follow_redirects=True)
    page = resp.get_data(as_text=True)
    assert "다른 사람이 먼저 수정" in page  # 비고 충돌 경고
    detail = b.get(f"/task/{tid}").get_data(as_text=True)
    assert "A 비고" in detail and "B 비고" not in detail
    assert "2026-10-01" in detail  # B의 완료일은 반영


def test_bulk_done(client_for):
    c = client_for("이가람C", "pc1")
    ids = [_task_id(c, 1), _task_id(c, 2)]
    c.post("/bulk_done", data={"ids": ids, "done_date": "2026-09-30"})
    html = c.get("/?status=완료").get_data(as_text=True)
    assert html.count("badge b-done") == 3  # 기존 완료 1건(No.47) + 2건


def test_unchanged_save_writes_nothing(client_for, tmp_path):
    """수정 화면에서 아무것도 안 바꾸고 저장하면 이력이 생기지 않아야 한다."""
    c = client_for("이가람C", "pc1")
    tid = _task_id(c, 6)  # 협업자·목표 이력이 있는 업무
    html = c.get(f"/task/{tid}/edit").get_data(as_text=True)
    orig = json.loads(unescape(html.split('name="orig" value="', 1)[1].split('"', 1)[0]))
    data = {k: (v or "") for k, v in orig.items() if k != "collaborators"}
    data["collaborators"] = orig["collaborators"]
    data["orig"] = json.dumps(orig)
    c.post(f"/task/{tid}/edit", data=data)
    store = Store(tmp_path)
    store.refresh()
    assert [a for a in store.tasks[tid]["audit"] if a["by"] != "이관"] == []


def test_cross_site_post_blocked(client_for):
    c = client_for("나대경M", "pc1")
    resp = c.post("/bulk_done", data={"ids": []}, headers={"Origin": "http://evil.example"})
    assert resp.status_code == 403
